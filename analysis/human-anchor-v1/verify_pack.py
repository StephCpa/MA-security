"""Read-only verification of exported human-anchor workbooks."""
import hashlib
import json
from pathlib import Path
from collections import Counter
import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[3] / 'outputs/01a000ff-9c62-7890-8c26-b78b5d02aab6-human-anchor'
KEY = json.loads((HERE/'coordinator_only/key.json').read_text(encoding='utf8'))
checks=[]
def check(name, condition):
    checks.append({'name':name,'pass':bool(condition)})
for rater in ('A','B'):
    wb=openpyxl.load_workbook(OUT/f'F7_human_{rater}.xlsx')
    check('sheet_names_'+rater,wb.sheetnames==['Ratings','Practice','Instructions'])
    check('no_hidden_sheets_'+rater,all(s.sheet_state=='visible' for s in wb))
    for name,kind,n in [('Ratings','main',96),('Practice','practice',6)]:
        ws=wb[name]
        key=sorted((k for k in KEY if k['rater']==rater and k['kind']==kind),key=lambda k:k['position'])
        check(f'{rater}_{name}_count',ws.max_row==n+1 and ws.max_column==7 and len(key)==n)
        check(f'{rater}_{name}_validations',len(ws.data_validations.dataValidation)==2)
        check(f'{rater}_{name}_freeze',ws.freeze_panes=='A2')
        for row,k in zip(ws.iter_rows(min_row=2),key):
            check('text_'+k['item_id'],row[0].value==k['item_id'] and hashlib.sha256(row[1].value.encode()).hexdigest()==k['output_sha256'])
            check('blank_'+k['item_id'],all(c.value is None for c in row[2:]))
        check(f'{rater}_{name}_no_formulas',not any(c.data_type=='f' for row in ws for c in row))
    main=[k for k in KEY if k['rater']==rater and k['kind']=='main']
    counts=Counter(k['pair_id'] for k in main)
    check('pairs_'+rater,len(counts)==48 and all(v==2 for v in counts.values()))
    distances=[abs(a['position']-b['position']) for i,a in enumerate(main) for b in main[i+1:] if a['pair_id']==b['pair_id']]
    check('pair_spacing_'+rater,len(distances)==48 and min(distances)>=6)
check('no_shared_ids',len({k['item_id'] for k in KEY})==204)
main_a=[k for k in KEY if k['rater']=='A' and k['kind']=='main']
main_b=[k for k in KEY if k['rater']=='B' and k['kind']=='main']
check('same_content_different_order',
      {(k['pair_id'],k['form'],k['output_sha256']) for k in main_a} == {(k['pair_id'],k['form'],k['output_sha256']) for k in main_b}
      and [k['output_sha256'] for k in main_a] != [k['output_sha256'] for k in main_b])
inspection=HERE.parents[2]/'.artifact-work/five-form-confirmation-ratings-20260828/confirmation_scored_inspection.json'
old=json.loads(inspection.read_text(encoding='utf8'))
old_keys={(k['pair_id'],k['generated_form'],k['output_sha256']) for k in old['key']}
check('matches_original_model_rated_texts',old_keys=={(k['pair_id'],k['form'],k['output_sha256']) for k in main_a})
result={'checks':checks,'passed':all(c['pass'] for c in checks),'check_count':len(checks),'ratings_collected':0}
(HERE/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in result.items() if k!='checks'}))
if not result['passed']:raise SystemExit(1)
