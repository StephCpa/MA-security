"""Read-only human return validation and source-pair clustered agreement."""
import argparse
import hashlib
import json
import math
import platform
from pathlib import Path

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
FIELDS = ['retention_0_3', 'centrality_0_3', 'provenance_0_3', 'relay_mode']
MODEL_FIELDS = ['actionable_semantic_retention', 'pragmatic_instruction_centrality', 'provenance_signal_strength']
MODES = ['IMP', 'REP', 'DROP', 'MIX']
HEADER = ['item_id', 'candidate_text', *FIELDS, 'reason_or_uncertainty']
REPS = 5000
SEED = 2026092807


def validate_rows(rows, key):
    expected = {k['item_id']: k for k in key}
    if len(rows) != len(key):
        raise ValueError('row count changed')
    seen, result = set(), {}
    for row in rows:
        if len(row) != 7:
            raise ValueError('column count changed')
        item, text, *rest = row
        if item not in expected or item in seen:
            raise ValueError('unknown or duplicate item id')
        seen.add(item)
        k = expected[item]
        if not isinstance(text, str) or hashlib.sha256(text.encode()).hexdigest() != k['output_sha256']:
            raise ValueError('candidate text changed')
        scores = []
        for value in rest[:3]:
            if value is None or value == '':
                scores.append(None)
            elif type(value) in (int, float) and math.isfinite(value) and int(value) == value and 0 <= value <= 3:
                scores.append(int(value))
            else:
                raise ValueError('ordinal score must be numeric integer 0..3 or blank')
        mode = rest[3]
        if mode is None or mode == '':
            mode = None
        elif mode not in MODES:
            raise ValueError('unknown relay-mode code')
        note = rest[4]
        if note is not None and not isinstance(note, str):
            raise ValueError('reason must be text')
        result[k['pair_id'], k['form']] = dict(zip(FIELDS, [*scores, mode]), reason=note)
    return result


def read_return(filename, rater, all_key):
    book = openpyxl.load_workbook(filename, data_only=False)
    if book.sheetnames != ['Ratings', 'Practice', 'Instructions']:
        raise ValueError('unexpected sheet layout')
    sheet = book['Ratings']
    if [c.value for c in sheet[1]] != HEADER:
        raise ValueError('header changed')
    if any(c.data_type == 'f' for row in sheet for c in row):
        raise ValueError('formula in returned rating data')
    key = [k for k in all_key if k['rater'] == rater and k['kind'] == 'main']
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    return validate_rows(rows, key)


def metrics(matrix, ordinal):
    n = matrix.sum()
    if n == 0:
        return {'exact_agreement': None, 'kappa': None}
    p = matrix / n
    weights = 1 - np.abs(np.arange(4)[:, None] - np.arange(4)[None, :]) / 3 if ordinal else np.eye(4)
    observed = float((p * weights).sum())
    chance = float((np.outer(p.sum(axis=1), p.sum(axis=0)) * weights).sum())
    return {'exact_agreement': float(np.trace(p)), 'kappa': (observed-chance)/(1-chance) if 1-chance > 1e-12 else None}


def comparison(left, right, field, reps=REPS):
    keys = sorted(set(left) | set(right))
    pairs = sorted({key[0] for key in keys})
    matrices = np.zeros((len(pairs), 4, 4), dtype=int)
    pair_index = {p: i for i, p in enumerate(pairs)}
    missing = {'left_only_missing': 0, 'right_only_missing': 0, 'both_missing': 0}
    for key in keys:
        a, b = left.get(key, {}).get(field), right.get(key, {}).get(field)
        if a is None or b is None:
            missing['both_missing' if a is None and b is None else 'left_only_missing' if a is None else 'right_only_missing'] += 1
            continue
        if field == 'relay_mode':
            a, b = MODES.index(a), MODES.index(b)
        matrices[pair_index[key[0]], a, b] += 1
    total = matrices.sum(axis=0)
    point = metrics(total, field != 'relay_mode')
    rng = np.random.default_rng(SEED)
    strata = {}
    for index, pair in enumerate(pairs):
        # Confirmation IDs are cf-FAMILY-NN. Preserve the six authored strata.
        strata.setdefault(pair.rsplit('-', 1)[0], []).append(index)
    samples = {name: [] for name in point}
    for _ in range(reps):
        indices = np.concatenate([rng.choice(group, size=len(group), replace=True) for group in strata.values()])
        sample = matrices[indices].sum(axis=0)
        measured = metrics(sample, field != 'relay_mode')
        for name, value in measured.items():
            if value is not None:
                samples[name].append(value)
    intervals = {}
    for name, values in samples.items():
        intervals[name] = {'valid_draws': len(values), 'undefined_draws': reps-len(values),
                           'percentile_95': np.quantile(values, [.025, .975]).tolist() if len(values) == reps and point[name] is not None else None}
    return {'source_pairs': len(pairs), 'rated_text_pairs': int(total.sum()), 'missing': missing,
            'categories': list(range(4)) if field != 'relay_mode' else MODES,
            'confusion_left_rows_right_columns': total.tolist(), 'point': point,
            'bootstrap': intervals, 'draws': reps, 'seed': SEED,
            'stratum_pair_counts': {k: len(v) for k,v in strata.items()},
            'weighting': 'linear' if field != 'relay_mode' else 'unweighted_nominal'}


def old_model_ratings():
    source = HERE.parents[2] / '.artifact-work/five-form-confirmation-ratings-20260828/confirmation_scored_inspection.json'
    old = json.loads(source.read_text(encoding='utf8'))
    out = {}
    for rater in ('A', 'B'):
        keys = sorted([k for k in old['key'] if k['rater'] == rater], key=lambda k: int(k['position']))
        scores = old['scores'][rater]
        assert len(keys) == len(scores) == 96
        values = {}
        for k, score in zip(keys, scores):
            assert k['item_id'] == score['item_id']
            assert hashlib.sha256(score['candidate_text'].encode()).hexdigest() == k['output_sha256']
            values[k['pair_id'], k['generated_form']] = {new: score[old_name] for new, old_name in zip(FIELDS, MODEL_FIELDS)}
        out['model_' + rater] = values
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--a', required=True)
    parser.add_argument('--b', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--validate-only', action='store_true')
    parser.add_argument('--attestation')
    args = parser.parse_args()
    key = json.loads((HERE/'coordinator_only/key.json').read_text(encoding='utf8'))
    humans = {r: read_return(getattr(args, r.lower()), r, key) for r in ('A', 'B')}
    result = {'status': 'validation_only', 'files_sha256': {r: hashlib.sha256(Path(getattr(args, r.lower())).read_bytes()).hexdigest() for r in humans},
              'software': {'python': platform.python_version(), 'numpy': np.__version__, 'openpyxl': openpyxl.__version__},
              'missing_by_rater': {r: {f: sum(v[f] is None for v in data.values()) for f in FIELDS} for r, data in humans.items()},
              'unexplained_missing_item_ids': {r: [k['item_id'] for k in key if k['rater']==r and k['kind']=='main' and any(humans[r][k['pair_id'],k['form']][f] is None for f in FIELDS) and not humans[r][k['pair_id'],k['form']]['reason']] for r in humans}}
    if not args.validate_only:
        frozen = json.loads((HERE/'coordinator_only/analysis_freeze.json').read_text(encoding='utf8'))
        for relative, digest in frozen['file_sha256'].items():
            if hashlib.sha256((HERE/relative).read_bytes()).hexdigest() != digest:
                raise ValueError('frozen analysis dependency changed: ' + relative)
        if not args.attestation:
            raise ValueError('coordinator human/blinding attestation required; not inferred from filenames')
        attestation = json.loads(Path(args.attestation).read_text(encoding='utf8'))
        for r in humans:
            if any(attestation.get(r, {}).get(f) is not True for f in ('human', 'no_model_assistance', 'no_main_cross_visibility', 'outcome_blind')):
                raise ValueError('independent blind-human claim not supported; preserve returns for a separately labeled deviation analysis')
            if all(v[f] is None for v in humans[r].values() for f in FIELDS):
                raise ValueError('blank workbook is not a human return')
        result['attestation'] = attestation
        models = old_model_ratings()
        if any(set(data) != set(humans['A']) for data in models.values()) or set(humans['A']) != set(humans['B']):
            raise ValueError('rating universes differ')
        results = {'human_A_vs_human_B': {f: comparison(humans['A'], humans['B'], f) for f in FIELDS}}
        for r, data in humans.items():
            for m, model in models.items():
                results[f'human_{r}_vs_{m}'] = {f: comparison(data, model, f) for f in FIELDS[:3]}
        result.update(status='human_agreement_analysis', comparisons=results,
                      scope='Exploratory agreement within 48 authored source pairs; pointwise intervals, no tests or causal shares. Unknown kappa remains null. Practice excluded.')
    with Path(args.output).open('x', encoding='utf8') as stream:
        json.dump(result, stream, indent=2, allow_nan=False)
        stream.write('\n')
    print(json.dumps({'status': result['status'], 'missing_by_rater': result['missing_by_rater']}))


if __name__ == '__main__':
    main()
